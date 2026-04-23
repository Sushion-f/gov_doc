import hashlib
import json
import shutil
from datetime import datetime
from html import unescape
from io import BytesIO
from pathlib import Path, PurePosixPath
import re

from .config import settings


def data_root() -> Path:
    root = Path(settings.workspace_root)
    try:
        root.mkdir(parents=True, exist_ok=True)
        return root
    except OSError:
        fallback = Path(__file__).resolve().parent.parent / ".runtime-data" / "users"
        fallback.mkdir(parents=True, exist_ok=True)
        return fallback


def user_root(user_id: str) -> Path:
    root = data_root() / user_id
    root.mkdir(parents=True, exist_ok=True)
    return root


def user_workspace_root(user_id: str) -> Path:
    root = user_root(user_id) / "workspace"
    root.mkdir(parents=True, exist_ok=True)
    return root


def workspace_files_root(user_id: str) -> Path:
    root = user_workspace_root(user_id) / "files"
    root.mkdir(parents=True, exist_ok=True)
    return root


def workspace_cache_root(user_id: str) -> Path:
    root = user_workspace_root(user_id) / ".cache"
    root.mkdir(parents=True, exist_ok=True)
    return root


def workspace_versions_root(user_id: str) -> Path:
    root = user_workspace_root(user_id) / ".versions"
    root.mkdir(parents=True, exist_ok=True)
    return root


def workspace_trash_root(user_id: str) -> Path:
    root = user_workspace_root(user_id) / ".trash"
    root.mkdir(parents=True, exist_ok=True)
    return root


def user_memory_root(user_id: str) -> Path:
    root = user_root(user_id) / "memory"
    root.mkdir(parents=True, exist_ok=True)
    return root


def user_template_root(user_id: str) -> Path:
    root = user_root(user_id) / "templates"
    root.mkdir(parents=True, exist_ok=True)
    return root


def normalize_workspace_path(relative_path: str | None) -> str:
    raw = str(relative_path or "").replace("\\", "/").strip().strip("/")
    if not raw:
        return ""
    normalized = PurePosixPath(raw)
    parts = [part for part in normalized.parts if part not in {"", "."}]
    if any(part == ".." for part in parts):
        raise ValueError("非法工作区路径")
    return "/".join(parts)


def workspace_abspath(user_id: str, relative_path: str | None) -> Path:
    normalized = normalize_workspace_path(relative_path)
    root = workspace_files_root(user_id)
    return root if not normalized else root / normalized


def workspace_relpath(user_id: str, absolute_path: Path) -> str:
    return normalize_workspace_path(str(absolute_path.relative_to(workspace_files_root(user_id))))


def read_user_doc(user_id: str, relative_path: str | None) -> str:
    if not relative_path:
        return ""
    path = user_root(user_id) / relative_path
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def clear_user_memory_subdir(user_id: str, relative_dir: str) -> None:
    path = user_memory_root(user_id) / relative_dir
    if path.exists():
        shutil.rmtree(path)


def write_memory_doc(user_id: str, name: str, content: str) -> str:
    root = user_memory_root(user_id)
    path = root / name
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        existing = path.read_text(encoding="utf-8")
        if existing == content:
            return str(path.relative_to(user_root(user_id)))
    path.write_text(content, encoding="utf-8")
    return str(path.relative_to(user_root(user_id)))


def write_template_file(user_id: str, file_name: str, content: bytes) -> str:
    root = user_template_root(user_id)
    path = root / file_name
    path.write_bytes(content)
    return str(path.relative_to(user_root(user_id)))


def workspace_file_etag(file_path: Path) -> str:
    stat = file_path.stat()
    payload = f"{file_path.resolve()}:{stat.st_mtime_ns}:{stat.st_size}"
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()


def workspace_cache_key(file_path: Path) -> str:
    stat = file_path.stat()
    payload = f"{file_path.resolve()}:{stat.st_mtime_ns}"
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()


def workspace_cache_dir_for(user_id: str, relative_path: str) -> Path:
    file_path = workspace_abspath(user_id, relative_path)
    cache_dir = workspace_cache_root(user_id) / workspace_cache_key(file_path)
    cache_dir.mkdir(parents=True, exist_ok=True)
    return cache_dir


def write_parse_cache(
    user_id: str,
    relative_path: str,
    *,
    content_text: str,
    content_html: str,
    meta: dict,
) -> dict:
    cache_dir = workspace_cache_dir_for(user_id, relative_path)
    (cache_dir / "content.txt").write_text(content_text or "", encoding="utf-8")
    (cache_dir / "content.html").write_text(content_html or "<pre></pre>", encoding="utf-8")
    meta_path = cache_dir / "meta.json"
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "content_text": content_text or "",
        "content_html": content_html or "<pre></pre>",
        "meta": meta,
        "cache_dir": cache_dir,
    }


def read_parse_cache(user_id: str, relative_path: str) -> dict | None:
    file_path = workspace_abspath(user_id, relative_path)
    if not file_path.exists() or file_path.is_dir():
        return None
    cache_dir = workspace_cache_root(user_id) / workspace_cache_key(file_path)
    meta_path = cache_dir / "meta.json"
    text_path = cache_dir / "content.txt"
    html_path = cache_dir / "content.html"
    if not (meta_path.exists() and text_path.exists() and html_path.exists()):
        return None
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None
    if meta.get("etag") != workspace_file_etag(file_path):
        return None
    # 如果上次解析走了 fallback 路径（依赖库缺失时的裸字节兜底），
    # 现在依赖库可能已经补齐，应该丢弃缓存、重新解析一次。
    parser_name = str(meta.get("parser") or "")
    if parser_name.endswith("-fallback"):
        return None
    return {
        "content_text": text_path.read_text(encoding="utf-8"),
        "content_html": html_path.read_text(encoding="utf-8"),
        "meta": meta,
        "cache_dir": cache_dir,
    }


def parse_workspace_file(user_id: str, relative_path: str, *, force: bool = False) -> dict:
    normalized = normalize_workspace_path(relative_path)
    if not normalized:
        raise FileNotFoundError("工作区路径为空")
    file_path = workspace_abspath(user_id, normalized)
    if not file_path.exists() or file_path.is_dir():
        raise FileNotFoundError(f"文件不存在: {normalized}")
    if not force:
        cached = read_parse_cache(user_id, normalized)
        if cached is not None:
            return cached

    from .file_parser import parse_file_path

    parsed = parse_file_path(file_path)
    meta = {
        "etag": workspace_file_etag(file_path),
        "parsedAt": datetime.utcnow().isoformat(),
        "parser": parsed["meta"].get("parser", "unknown"),
        "pages": parsed["meta"].get("pages"),
        "path": normalized,
        "size": file_path.stat().st_size,
    }
    meta.update({k: v for k, v in parsed["meta"].items() if k not in meta})
    return write_parse_cache(
        user_id,
        normalized,
        content_text=parsed["content_text"],
        content_html=parsed["content_html"],
        meta=meta,
    )


def ensure_unique_workspace_path(
    user_id: str,
    desired_relative_path: str,
    *,
    is_dir: bool,
    exclude_relative_path: str | None = None,
) -> str:
    normalized = normalize_workspace_path(desired_relative_path)
    if not normalized:
        raise ValueError("工作区路径不能为空")
    exclude_normalized = normalize_workspace_path(exclude_relative_path) if exclude_relative_path else ""
    candidate = PurePosixPath(normalized)
    parent = "" if str(candidate.parent) == "." else str(candidate.parent)
    stem = candidate.stem or candidate.name
    suffix = candidate.suffix
    index = 1
    while True:
        current = normalize_workspace_path(str(candidate))
        if current == exclude_normalized:
            return current
        absolute = workspace_abspath(user_id, current)
        if not absolute.exists():
            return current
        name = candidate.name if is_dir else f"{stem} ({index}){suffix}"
        if is_dir:
            name = f"{candidate.name} ({index})"
        candidate = PurePosixPath(parent, name) if parent else PurePosixPath(name)
        index += 1


def write_workspace_bytes(user_id: str, relative_path: str, content: bytes) -> Path:
    path = workspace_abspath(user_id, relative_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return path


def move_workspace_path(user_id: str, old_relative_path: str, new_relative_path: str) -> Path:
    source = workspace_abspath(user_id, old_relative_path)
    target = workspace_abspath(user_id, new_relative_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    if source.exists():
        source.rename(target)
    return target


def trash_workspace_path(user_id: str, relative_path: str) -> Path | None:
    source = workspace_abspath(user_id, relative_path)
    if not source.exists():
        return None
    trash_name = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{source.name}"
    target = workspace_trash_root(user_id) / trash_name
    source.rename(target)
    return target


def html_to_plain_text(content_html: str | None, fallback_text: str | None = None) -> str:
    html = str(content_html or "")
    text = re.sub(r"(?i)<br\s*/?>", "\n", html)
    text = re.sub(r"(?i)</p\s*>", "\n\n", text)
    text = re.sub(r"(?i)</div\s*>", "\n", text)
    text = re.sub(r"(?i)<li\s*>", "- ", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = unescape(text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if text:
        return _sanitize_xml_text(text)
    return _sanitize_xml_text((fallback_text or "").strip())


def _sanitize_xml_text(value: str) -> str:
    return "".join(
        ch
        for ch in str(value or "")
        if ch in {"\t", "\n", "\r"} or ord(ch) >= 32
    )


def build_docx_bytes(title: str, content_html: str | None, content_text: str | None) -> bytes:
    from docx import Document

    document = Document()
    paragraphs = [
        _sanitize_xml_text(item.strip())
        for item in html_to_plain_text(content_html, content_text).splitlines()
    ]
    if not any(paragraphs):
        document.add_paragraph("")
    else:
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def resolve_writable_file_path(relative_path: str, *, create_companion_for_binary: bool = True) -> str:
    normalized = normalize_workspace_path(relative_path)
    ext = Path(normalized).suffix.lower()
    if ext in {".txt", ".md", ".markdown", ".json", ".csv", ".docx", ""}:
        return normalized
    if not create_companion_for_binary:
        return normalized
    path_obj = PurePosixPath(normalized)
    companion_name = f"{path_obj.stem}.edit.docx"
    parent = "" if str(path_obj.parent) == "." else str(path_obj.parent)
    return normalize_workspace_path(str(PurePosixPath(parent, companion_name) if parent else PurePosixPath(companion_name)))


def write_workspace_version(
    user_id: str,
    file_id: str,
    version_no: int,
    title: str,
    content_html: str | None,
    content_text: str | None,
    annotations: list[dict] | None = None,
    *,
    file_relative_path: str | None = None,
    extra_meta: dict | None = None,
) -> str:
    version_dir = workspace_versions_root(user_id) / file_id / f"v{version_no:04d}"
    version_dir.mkdir(parents=True, exist_ok=True)
    (version_dir / "content.txt").write_text(content_text or "", encoding="utf-8")
    meta = {
        "title": title or "",
        "annotations": annotations or [],
        "filePath": normalize_workspace_path(file_relative_path),
        "savedAt": datetime.utcnow().isoformat(),
        "hasHtml": bool(content_html),
    }
    if extra_meta:
        meta.update(extra_meta)
    (version_dir / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(version_dir.relative_to(user_root(user_id)))


def latest_workspace_version(user_id: str, file_id: str) -> dict | None:
    root = workspace_versions_root(user_id) / file_id
    if not root.exists():
        return None
    versions = sorted([item for item in root.iterdir() if item.is_dir()], key=lambda item: item.name)
    if not versions:
        return None
    latest = versions[-1]
    meta_path = latest / "meta.json"
    content_path = latest / "content.txt"
    meta = {}
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            meta = {}
    return {
        "version": latest.name,
        "content_text": content_path.read_text(encoding="utf-8") if content_path.exists() else "",
        "meta": meta,
        "path": latest,
    }

